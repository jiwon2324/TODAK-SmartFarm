# -*- coding: utf-8 -*-
"""표지 공통: 학교 로고·작성자 (PPT)"""
import os

from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "docs", "brand", "hywoman_logo.png")
STUDENT_LINE = "한양여자대학교 · 2402180 양지원  |  데이터 기획 및 분석 · 1학기 중간고사 대체과제"


def student_badge(sl):
    """표지: 오른쪽 위 흰 카드에 학교 로고, 아래에 작성자"""
    from pptx.enum.shapes import MSO_SHAPE
    card = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(0.45), Inches(3.95), Inches(1.05))
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF); card.line.fill.background()
    card.adjustments[0] = 0.12
    if os.path.exists(LOGO):
        sl.shapes.add_picture(LOGO, Inches(9.1), Inches(0.6), width=Inches(3.55))  # 3.55 x 0.74 (원본 비율 4.8:1)
    tb = sl.shapes.add_textbox(Inches(0.8), Inches(5.55), Inches(11.5), Inches(0.5)).text_frame
    tb.word_wrap = True
    r = tb.paragraphs[0].add_run(); r.text = STUDENT_LINE
    r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


TODAK_WHITE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "docs", "brand", "todak_logo_white.png")


def todak_logo(sl, x=0.8, y=1.15, h=0.85):
    """표지(어두운 배경): 흰색 TODAK 로고"""
    if os.path.exists(TODAK_WHITE):
        sl.shapes.add_picture(TODAK_WHITE, Inches(x), Inches(y), height=Inches(h))
        return True
    return False
