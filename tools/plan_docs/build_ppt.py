# -*- coding: utf-8 -*-
"""산출물 1 발표 자료: 시스템 개발 계획서 PPT (python-pptx, 네이티브 차트)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
import data as D
import spec as S

FONT = "맑은 고딕"
FIG = os.path.join(os.path.dirname(__file__), "fig")
OUT = os.path.join(os.path.dirname(__file__), "..", "TODAK_산출물", "02_시스템_개발_계획서_v2.0.pptx")

DARK = RGBColor(0x1B, 0x3D, 0x20); GREEN = RGBColor(0x2E, 0x7D, 0x32); MOSS = RGBColor(0xA5, 0xD6, 0xA7)
TINT = RGBColor(0xEE, 0xF6, 0xEE); INK = RGBColor(0x20, 0x21, 0x24); MUTED = RGBColor(0x5F, 0x63, 0x68)
ORANGE = RGBColor(0xE6, 0x51, 0x00); WHITE = RGBColor(0xFF, 0xFF, 0xFF); LINE = RGBColor(0xD6, 0xDB, 0xD6)
BLUE = RGBColor(0x15, 0x65, 0xC0)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5
page = [0]


def tf_text(tf, text, size=14, bold=False, color=INK, align=None):
    tf.clear()
    lines = text if isinstance(text, list) else [text]
    for i, l in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        parts = l.split("**")
        for k, t in enumerate(parts):
            if not t:
                continue
            r = p.add_run(); r.text = t
            r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold or (k % 2 == 1); r.font.color.rgb = color
            r._r.get_or_add_rPr().set("lang", "ko-KR")
            ea = r._r.get_or_add_rPr().find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
            if ea is None:
                from lxml import etree
                ea = etree.SubElement(r._r.get_or_add_rPr(), "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
            ea.set("typeface", FONT)
        p.space_after = Pt(6)


def text(sl, x, y, w, h, t, size=14, bold=False, color=INK, align=None, anchor=MSO_ANCHOR.TOP, name=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    tf_text(tf, t, size, bold, color, align)
    if name:
        tb.name = name
    return tb


def card(sl, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line; s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def base(title, sub=None, dark=False, notes=""):
    sl = prs.slides.add_slide(BLANK)
    page[0] += 1
    bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = DARK if dark else WHITE
    if not dark:
        text(sl, 0.6, 0.38, 12.1, 0.8, title, 28, True, INK, name="Title")
        if sub:
            text(sl, 0.6, 1.1, 12.1, 0.45, sub, 14, False, MUTED, name="Subtitle")
        text(sl, 0.6, 7.02, 8, 0.3, "TODAK 시스템 개발 계획서 v2.0  |  양평 100평 방울토마토 스마트 농업 정보시스템", 10, False, MUTED)
        text(sl, 11.9, 7.02, 0.9, 0.3, str(page[0]), 10, False, MUTED, PP_ALIGN.RIGHT)
    if notes:
        sl.notes_slide.notes_text_frame.text = notes
    return sl


def pic(sl, name, x, y, w=None, h=None):
    path = os.path.join(FIG, name + ".png")
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return sl.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def fit_pic(sl, name, x, y, maxw, maxh):
    from PIL import Image
    im = Image.open(os.path.join(FIG, name + ".png")); r = im.width / im.height
    w = maxw; h = w / r
    if h > maxh:
        h = maxh; w = h * r
    return pic(sl, name, x + (maxw - w) / 2, y, w=w)


def table(sl, x, y, w, headers, rows, widths, size=11, rowh=0.36):
    nr, nc = len(rows) + 1, len(headers)
    gt = sl.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr)).table
    tot = sum(widths)
    for i, cw in enumerate(widths):
        gt.columns[i].width = Inches(w * cw / tot)
    for r in range(nr):
        gt.rows[r].height = Inches(rowh)
        for c in range(nc):
            cell = gt.cell(r, c)
            val = headers[c] if r == 0 else rows[r - 1][c]
            cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf_text(cell.text_frame, str(val), size, r == 0, WHITE if r == 0 else INK)
            for p in cell.text_frame.paragraphs:
                p.space_after = Pt(0)
            cell.fill.solid(); cell.fill.fore_color.rgb = GREEN if r == 0 else (TINT if r % 2 == 0 else WHITE)
    return gt


def style_chart(ch, legend=True):
    ch.font.name = FONT; ch.font.size = Pt(11); ch.font.color.rgb = MUTED
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(11)
    try:
        ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
        ch.value_axis.format.line.fill.background()
        ch.category_axis.format.line.color.rgb = LINE
    except Exception:
        pass


# 1 표지
sl = base("", dark=True, notes="TODAK은 경기도 양평군 100평 비닐하우스에서 방울토마토를 키우는 농가를 위한 웹 기반 스마트 농업 정보시스템입니다. 오늘은 개발 착수 전 계획을 보고드리고, 요구사항과 설계, 일정·예산을 승인받는 것이 목적입니다.")
text(sl, 0.8, 1.6, 6, 0.5, "TODAK", 20, True, MOSS)
text(sl, 0.8, 2.2, 11.5, 1.2, "방울토마토 스마트 농업 정보시스템\n시스템 개발 계획서", 40, True, WHITE)
text(sl, 0.8, 4.3, 11, 0.5, "경기도 양평군 · 비닐하우스 1동 100평 · 사업비 5,000만 원 · 12주 · 반응형 웹", 18, False, MOSS)
text(sl, 0.8, 6.3, 11, 0.4, "TODAK-PLN-001  v2.0  |  2026.10.26  |  검토용 (개발 전)", 12, False, MOSS)

# 2 요약
sl = base("한 장 요약: 양평 100평 하우스를 지키는 웹 시스템", notes="숫자 세 개만 기억해 주세요. 100평, 5천만 원, 12주입니다. 핵심 기능은 양평의 겨울 저온과 여름 고온을 감시하고 자동 제어하는 것이고, 장비가 없어도 시뮬레이터로 같은 코드를 검증할 수 있게 설계했습니다.")
for i, (n, l) in enumerate([("100평", "단동 비닐하우스 1동\n(330.58㎡, 2개 관리 구역)"), ("5,000만 원", "부가세 포함 고정\n(모두 추정 예산)"), ("12주", "2026.11.02 ~ 2027.01.22\n폭포수 · 관문 4개")]):
    card(sl, 0.6 + i * 4.1, 1.5, 3.8, 2.1)
    text(sl, 0.8 + i * 4.1, 1.65, 3.4, 0.9, n, 40, True, GREEN)
    text(sl, 0.8 + i * 4.1, 2.6, 3.4, 0.9, l, 13, False, MUTED)
msgs = ["**무엇을**: 온도·습도·CO₂·일사·배지수분·EC·pH 11채널을 1분마다 수집하고 환기·난방·차광·관수를 자동 제어",
        "**왜**: 양평 1월 평균 최저 −8.6℃, 8월 평균 최고 30.3℃ (기상청 평년값)",
        "**어떻게**: React 웹 + Spring Boot + MySQL 20개 테이블 + MQTT, 서버 판단·게이트웨이 안전 운전",
        "**검증**: 요구사항 25건(기능 18·비기능 7), 테스트 19건 — 모두 미실행, 시뮬레이터 검증과 실장비 검증을 구분"]
text(sl, 0.6, 4.0, 12.1, 2.8, msgs, 15)

# 3 개요
sl = base("교수 확정 조건과 사업 개요", "확정 조건은 바꾸지 않고, 확정되지 않은 시설 세부는 가정(A-01~07)으로 관리", notes="교수님이 지정한 조건은 지역, 작물, 시설, 면적, 사업비입니다. 하우스 크기나 재배 방식처럼 정해지지 않은 부분은 가정으로 표시하고 발주처 확인을 받겠습니다.")
table(sl, 0.6, 1.7, 6.2, ["항목", "내용"], [("지역", "경기도 양평군"), ("작물", "방울토마토"), ("시설", "비닐하우스 1동, 100평"), ("사업비", "50,000,000원 (부가세 포함)"),
                                          ("기간", "12주 (문서상 가상 기간)"), ("개발", "반응형 웹, 폭포수")], [1.4, 4], 13, 0.45)
table(sl, 7.1, 1.7, 5.6, ["가정", "내용"], [(a[0], a[2]) for a in D.ASSUMPTIONS], [0.9, 4.7], 10.5, 0.55)

# 4 검토 결과
sl = base("기존 계획(v1.x) 검토: 다섯 가지를 고쳤다", notes="이전 문서에는 측정하지 않은 수치가 실적처럼 들어가 있었고, 누가 제어를 판단하는지가 두 가지로 적혀 있었습니다. 이 다섯 가지를 고쳐 v2.0을 만들었습니다.")
items = [("시설", "1,300평 연동형 → 100평 단동", "교수 조건 반영"), ("플랫폼", "Flutter 앱 → 반응형 웹", "12주 범위 현실화"),
         ("자동화 범위", "파종 '자동' → 범위 밖", "육묘 관리 ≠ 파종"), ("제어 주체", "두 가지 서술 → 서버 판단·GW 안전 운전", "설계 일관성"),
         ("가상 실적", "88%·2.1초·20건 통과 → 측정 대기", "검증 원칙")]
for i, (k, a, b) in enumerate(items):
    x = 0.6 + (i % 3) * 4.1; y = 1.5 + (i // 3) * 2.6
    card(sl, x, y, 3.8, 2.3)
    text(sl, x + 0.2, y + 0.15, 3.4, 0.4, k, 13, True, GREEN)
    text(sl, x + 0.2, y + 0.6, 3.4, 1.0, a, 16, True, INK)
    text(sl, x + 0.2, y + 1.65, 3.4, 0.5, b, 12, False, MUTED)
card(sl, 8.8, 4.1, 3.8, 2.3, fill=DARK)
text(sl, 9.0, 4.3, 3.4, 1.9, ["DB 9개 → 20개", "API 0개 → 27개", "출처 1건 → 11건"], 16, True, WHITE)

# 5 기후 기온
sl = base("양평은 11~3월 평균 최저가 영하, 8월 평균 최고는 30℃", "외부 기후이며 하우스 내부 값이 아님 · 기상청 평년값(1991~2020), 양평관측소", notes="그래프의 초록 선이 평균기온입니다. 겨울에는 평균 최저가 영하 8.6도까지 내려가고, 여름 최고는 30도를 넘습니다. 그래서 TODAK의 첫 번째 임무는 겨울 야간 난방 감시, 두 번째는 여름 환기입니다.")
cd = CategoryChartData(); cd.categories = D.MONTHS
cd.add_series("평균 최고(℃)", D.CLIMATE["tmax"]); cd.add_series("평균기온(℃)", D.CLIMATE["tavg"]); cd.add_series("평균 최저(℃)", D.CLIMATE["tmin"])
ch = sl.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.7), Inches(8.2), Inches(5.1), cd).chart
style_chart(ch)
for s, col in zip(ch.series, (ORANGE, GREEN, BLUE)):
    s.format.line.color.rgb = col; s.format.line.width = Pt(2.25); s.smooth = False
    s.marker.format.fill.solid(); s.marker.format.fill.fore_color.rgb = col; s.marker.format.line.color.rgb = col
card(sl, 9.1, 1.7, 3.6, 5.1)
text(sl, 9.3, 1.9, 3.2, 4.7, ["**−8.6℃**", "1월 평균 최저", "", "**30.3℃**", "8월 평균 최고", "", "**13℃ 이상**", "4·5·10월 일교차 → 새벽 난방과 한낮 환기가 하루에 함께"], 15)

# 6 강수·일조
sl = base("7~8월에 연 강수의 53%가 몰리고 7월 일조가 가장 적다", "다습·호우·광 부족이 같은 시기에 겹친다", notes="장마철인 7월에 비가 409밀리미터 오고 일조는 150시간으로 가장 적습니다. 다습하면 곰팡이병 위험이 커지므로 습도 경보와 환기 연동을 넣었습니다.")
cd = CategoryChartData(); cd.categories = D.MONTHS; cd.add_series("강수량(mm)", D.CLIMATE["rain"])
gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.7), Inches(6.1), Inches(5.0), cd)
ch = gf.chart; style_chart(ch, False); ch.has_title = True; ch.chart_title.text_frame.text = "월 강수량 (mm)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True; pl.data_labels.font.size = Pt(9); pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
for i, v in enumerate(D.CLIMATE["rain"]):
    pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = ORANGE if v >= 300 else GREEN
cd = CategoryChartData(); cd.categories = D.MONTHS; cd.add_series("일조시간(시간)", D.CLIMATE["sun"])
ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.9), Inches(1.7), Inches(5.8), Inches(5.0), cd).chart
style_chart(ch, False); ch.has_title = True; ch.chart_title.text_frame.text = "월 일조시간 (시간)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True; pl.data_labels.font.size = Pt(9); pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
for i, v in enumerate(D.CLIMATE["sun"]):
    pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = ORANGE if v < 160 else MOSS

# 7 기후→기능
sl = base("기후 자료가 요구사항과 테스트로 이어진다", "양평 기후 → 재배 위험 → 시스템 기능 → 검증", notes="기후 분석이 통계로 끝나지 않도록, 각 위험을 요구사항과 테스트 번호까지 연결했습니다.")
table(sl, 0.6, 1.7, 12.1, ["양평 기후", "재배 위험", "TODAK 기능", "검증"], [
    ("1월 평균최저 −8.6℃", "야간 저온·동해", "CR-02 야간 난방, CR-06 저온 안전, 오프라인 경보", "TC-002, TC-009"),
    ("8월 평균최고 30.3℃", "고온 낙화", "CR-01 환기, CR-03 차광, CR-05 고온 안전", "TC-002, TC-003"),
    ("7~8월 습도 79%", "곰팡이병", "습도 90% 초과 경보", "TC-009"),
    ("7월 일조 150.3시간", "광 부족", "DLI 계산, 부분 보광", "TC-003"),
    ("겨울 정전·통신 장애", "제어 중단", "UPS, 게이트웨이 안전 운전, 24시간 버퍼", "TC-N03")], [2.6, 2, 5.4, 2.1], 13, 0.62)

# 8 재배 기준
sl = base("재배 기준값은 '문헌값 → 초기 설정값 → 조정 범위'로 나눴다", "검색 요약 상태의 값은 네이버·원문 확인 후 확정 (과제 조건 미충족 항목)", notes="문헌값을 그대로 제어에 쓰지 않고, 시범운영 중 조정할 초기값을 따로 정했습니다. 과제 조건인 네이버 검색 확인은 아직 남아 있습니다.")
table(sl, 0.6, 1.7, 12.1, ["항목", "문헌값", "출처·검증", "TODAK 초기값"], [(g[0], g[1], f"{g[4]} · {g[5]}", g[6]) for g in D.GROW_STD[:8]],
      [2, 5.4, 2, 2.7], 11.5, 0.55)

# 9 GDD
sl = base("유효적산온도(GDD)는 내부·외부를 따로 계산한다", "수확일 예측이 아닌 생육 참고 지표", notes="외부 기온 기준으로는 겨울에 GDD가 전혀 쌓이지 않습니다. 하우스 내부 센서로 계산한 GDD와 나란히 보여 주면 난방이 생육에 얼마나 기여하는지 설명할 수 있습니다. 품종별 검증 모델이 없어 수확일은 예측하지 않습니다.")
card(sl, 0.6, 1.7, 5.4, 5.0)
text(sl, 0.8, 1.9, 5.0, 4.6, ["**GDD = max(0, min(Tavg, 30) − 10)**", "", "기준온도 10℃ · 상한 30℃ (Wikipedia GDD, 토마토 10℃)",
                               "", "**내부**: 센서 시간 평균 → 일평균", "**외부**: 기상청 (최고+최저)/2", "", "결측 20% 초과일은 '불완전' 표시", "기준온도 변경 이력 저장"], 14)
g = D.monthly_gdd_from_normals()
cd = CategoryChartData(); cd.categories = D.MONTHS; cd.add_series("외부 월 GDD (℃·일)", g)
ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.3), Inches(1.7), Inches(6.4), Inches(4.6), cd).chart
style_chart(ch, False); ch.has_title = True; ch.chart_title.text_frame.text = f"양평 외부 GDD 개략 (연 {sum(g):,.0f}℃·일)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
pl = ch.plots[0]; pl.gap_width = 60; pl.series[0].format.fill.solid(); pl.series[0].format.fill.fore_color.rgb = GREEN
text(sl, 6.3, 6.35, 6.4, 0.4, "평년 월평균기온으로 구한 개략치. 실제 구현은 일자료로 계산하며 내부 GDD는 측정 대기", 10, False, MUTED)

# 10 6가지 작업
sl = base("손 농사 6가지 중 완전 자동화는 온도 1개, 파종은 범위 밖", notes="여섯 가지 작업을 자동화 수준별로 나눴습니다. 온도는 완전 자동, 빛과 양액은 부분 자동, 벌 수정과 수확은 알림으로 사람을 돕고, 파종은 이번 범위에 넣지 않습니다.")
lv = {"완전 자동화(안전 한계 포함)": (GREEN, WHITE), "부분 자동화": (MOSS, INK), "의사결정 지원": (RGBColor(0xFF, 0xE0, 0xB2), INK), "수동 (범위 밖)": (RGBColor(0xEC, 0xEF, 0xF1), INK)}
for i, t in enumerate(D.MANUAL_TASKS):
    x = 0.6 + (i % 3) * 4.1; y = 1.5 + (i // 3) * 2.75
    card(sl, x, y, 3.8, 2.5)
    f, c = lv[t[3]]
    b = card(sl, x + 0.2, y + 0.2, 2.6, 0.42, fill=f)
    text(sl, x + 0.2, y + 0.2, 2.6, 0.42, t[3], 11, True, c, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(sl, x + 0.2, y + 0.75, 3.4, 0.45, t[0], 16, True)
    text(sl, x + 0.2, y + 1.25, 3.4, 1.2, t[4][:60] + ("…" if len(t[4]) > 60 else ""), 11.5, False, MUTED)

# 11 요구사항
sl = base("요구사항 25건: 기능 18 · 비기능 7, MVP 16건은 시뮬레이터로 완결", notes="기존 18건의 ID를 유지하고 구역 관리, GDD, 외부 기상, 명령 상태, 시뮬레이터, 실시간 갱신 6건을 추가했습니다. 오른쪽 매트릭스처럼 대부분 시뮬레이터만으로 개발과 검증이 가능합니다.")
for i, (n, l) in enumerate([("18", "기능 요구사항"), ("7", "비기능 요구사항"), ("6", "신규 (F-13~18)")]):
    card(sl, 0.6, 1.6 + i * 1.75, 3.4, 1.5)
    text(sl, 0.8, 1.7 + i * 1.75, 3.0, 0.8, n, 36, True, GREEN)
    text(sl, 0.8, 2.5 + i * 1.75, 3.0, 0.5, l, 13, False, MUTED)
fit_pic(sl, "f_req_matrix", 4.3, 1.6, 8.4, 5.2)

# 12 아키텍처
sl = base("서버가 판단하고, 현장 게이트웨이는 실행과 안전 운전을 맡는다", notes="전체 구조는 네 층입니다. 사용자는 브라우저로 접속하고, 서버는 Spring Boot와 MySQL, MQTT 브로커로 구성됩니다. 현장은 실장비를 쓰는 A안과 시뮬레이터를 쓰는 B안이 같은 토픽과 JSON을 쓰기 때문에 서버 코드는 같습니다.")
fit_pic(sl, "d_architecture", 0.6, 1.3, 12.1, 5.6)

# 13 제어 주체
sl = base("제어 주체를 하나로 정했다: 평상시 서버, 단절 시 게이트웨이", notes="기존 문서의 모순을 이렇게 정리했습니다. 우선순위는 안전 한계, 수동 명령, 자동 규칙, 단절 시 안전 운전 순서입니다.")
for i, (h, b, f) in enumerate([("평상시 (서버 연결)", ["판단: 서버 규칙 엔진, 1분마다", "실행: GW가 명령 수행 후 ACK", "기준값 원본: MySQL, 사본: GW", "수동 명령이 자동보다 우선"], TINT),
                               ("단절 시 (90초 무응답)", ["판단: GW 로컬 안전 로직", "범위: 고온 환기·저온 난방·관수 중지", "24시간 버퍼 후 재전송", "복구 시 기준값 버전 재동기화"], RGBColor(0xFF, 0xF3, 0xE0))]):
    card(sl, 0.6 + i * 6.2, 1.6, 5.9, 3.6, fill=f)
    text(sl, 0.9 + i * 6.2, 1.8, 5.4, 0.5, h, 18, True, GREEN if i == 0 else ORANGE)
    text(sl, 0.9 + i * 6.2, 2.5, 5.4, 2.6, b, 15)
card(sl, 0.6, 5.5, 12.1, 1.2, fill=DARK)
text(sl, 0.9, 5.6, 11.6, 1.0, "우선순위  ① 안전 한계(SAFETY)  →  ② 수동 명령(MANUAL)  →  ③ 자동 규칙(AUTO)  →  ④ 단절 시 안전 운전(FAILSAFE)", 14, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

# 14 통신
sl = base("통신 구조: 위로는 1분 측정값, 아래로는 명령과 ACK", notes="센서는 RS-485로 게이트웨이에 연결되고, 게이트웨이는 LTE와 MQTT로 서버에 1분 평균을 보냅니다. 명령은 반대 방향으로 내려가고 10초 안에 ACK가 없으면 한 번 재전송한 뒤 경보를 냅니다.")
fit_pic(sl, "d_comm_structure", 0.6, 1.3, 12.1, 5.6)

# 15 센서 배치
sl = base("100평을 2개 구역으로 나누고 구역마다 온습도·배지수분 1세트", "단동 8m × 41.5m (가정 A-01) · 환기·난방은 동 단위, 감시·관수는 구역 단위", notes="3구역은 41.5미터 길이에 비해 비용 대비 효과가 작고, 1구역은 센서가 고장 나면 대체할 수 없습니다. 그래서 2구역을 권장합니다.")
fit_pic(sl, "l_sensors", 0.6, 1.7, 12.1, 3.7)
table(sl, 0.6, 5.55, 12.1, ["구분", "수량", "근거"], [("온습도", "2세트", "구역별 대표 + 상호 검증"), ("배지수분", "2", "관수 계통별"), ("CO₂·일사·외부온도·EC·pH", "각 1", "단동 1개로 대표, 양액은 공급부")], [3, 1.5, 7.6], 11.5, 0.36)

# 16 장치 배치
sl = base("제어반에 게이트웨이·LTE·UPS를 모으고 간선을 길이 방향으로", notes="출입구 안쪽 제어반에 게이트웨이와 라우터, UPS를 모아 배선을 단순하게 했습니다. 난방기와 양액기는 농가 보유 설비를 연동한다는 가정입니다.")
fit_pic(sl, "l_devices", 0.6, 1.4, 12.1, 3.9)
text(sl, 0.6, 5.5, 12.1, 1.3, ["**기존 설비 연동 가정(A-06)**: 난방기·양액기·측창 모터·환기팬은 농가 보유 → 접점·릴레이 연동만", "확인 실패 시 위험 R-01 (가능성 3 × 영향 4) → 실장비 범위 축소"], 14)

# 17 제어 규칙
sl = base("자동 제어 규칙 6개: 히스테리시스와 안전 한계", notes="켜는 값과 끄는 값을 다르게 두고 최소 작동 시간을 줘서 임계값 근처에서 장치가 깜빡이지 않게 했습니다. 35도와 5도 안전 한계는 수동 모드보다 우선합니다.")
table(sl, 0.6, 1.6, 12.1, ["규칙", "이름", "ON", "OFF", "동작", "제한"], [(r[0], r[1], r[3], r[4], r[5], r[6]) for r in D.CONTROL_RULES], [1, 1.6, 2, 1.8, 4.2, 2.2], 12, 0.62)
text(sl, 0.6, 6.15, 12.1, 0.6, "실장비 적용 전: 시뮬레이터 시나리오 5종 통과 → 장치별 수동 시험 → 통신 차단 1시간 시험 → 첫 1주 '알림만' 모드", 13, False, MUTED)

# 18 DB
sl = base("데이터베이스 20개 테이블: 명령과 실제 상태를 분리했다", notes="가장 중요한 변경은 제어 명령과 실제 장치 상태를 다른 테이블로 나눈 것입니다. 그래야 실패하거나 시간 초과된 명령을 추적할 수 있습니다. 1분 데이터는 1년에 약 578만 건이 쌓입니다.")
fit_pic(sl, "u_erd", 0.6, 1.3, 8.6, 5.6)
for i, (n, l) in enumerate([("20", "테이블 (보완 9 + 신규 11)"), ("578만", "1분 데이터 행 / 년"), ("3년", "원시 데이터 보존")]):
    card(sl, 9.5, 1.4 + i * 1.85, 3.2, 1.6)
    text(sl, 9.7, 1.5 + i * 1.85, 2.9, 0.8, n, 30, True, GREEN)
    text(sl, 9.7, 2.3 + i * 1.85, 2.9, 0.5, l, 12, False, MUTED)

# 19 API
sl = base("REST API 27개 + MQTT 토픽 6개, 오류 코드까지 정의했다", notes="개발자가 추측 없이 구현할 수 있도록 경로, 권한, 요청과 응답 예시, 오류 코드를 계획서 18장에 모두 적었습니다. 여기서는 핵심만 보여 드립니다.")
key = [a for a in S.API if a[0] in ("API-01", "API-06", "API-07", "API-09", "API-12", "API-13", "API-16", "API-19", "API-21")]
table(sl, 0.6, 1.6, 12.1, ["ID", "메서드", "경로", "설명"], [(a[0], a[1], a[2], a[4]) for a in key], [1.1, 1.1, 4.6, 5.3], 11.5, 0.48)
text(sl, 0.6, 6.5, 12.1, 0.4, "오류: 400 · 401 · 403 · 404 · 409 · 423 SAFETY_LOCKED · 503 GATEWAY_OFFLINE · 500", 12, False, MUTED)

# 20 화면
sl = base("화면: 대시보드 하나에서 구역 상태·장치·알림을 본다", notes="PC에서는 왼쪽 메뉴와 카드 여덟 개, 스마트폰에서는 카드 한 줄과 하단 탭으로 같은 정보를 보여 줍니다. 화면의 숫자는 예시입니다.")
fit_pic(sl, "u_wireframe", 0.6, 1.3, 12.1, 5.3)
text(sl, 0.6, 6.65, 12.1, 0.3, "카드 값은 화면 예시이며 측정값이 아님", 10, False, MUTED)

# 21 일정
sl = base("12주 일정: 관문 4개에서 발주처 승인을 받는다", "문서상 가상 사업 기간 · 실제 학생 개발 일정은 별도로 확정", notes="2주차에 요구사항, 4주차에 설계, 10주차에 테스트, 12주차에 인수 승인을 받습니다. 구현과 현장 설치가 7주차부터 겹치는 것이 일정상 가장 빠듯한 부분입니다.")
fit_pic(sl, "f_gantt", 0.6, 1.7, 12.1, 5.2)

# 22 예산
tot = sum(b[2] for b in D.BUDGET)
sl = base("예산 5,000만 원: 인건비 50%, 현장 장비·설치 34%", "모두 견적 미확보 추정치 · 학생 실제 개발(B안) 실비와 섞지 않음", notes="예산은 고정 조건이라 합계 5천만 원 안에서 재배분했습니다. 대형 설비를 새로 사지 않고 농가 보유 설비에 연동한다는 가정이 깨지면 예산이 부족합니다.")
cd = CategoryChartData(); cd.categories = [b[0] for b in D.BUDGET][::-1]; cd.add_series("금액(백만 원)", [b[2] / 1e6 for b in D.BUDGET][::-1])
ch = sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.7), Inches(7.2), Inches(5.0), cd).chart
style_chart(ch, False); pl = ch.plots[0]; pl.gap_width = 50; pl.has_data_labels = True
pl.data_labels.number_format = '0.0"백만"'; pl.data_labels.number_format_is_linked = False; pl.data_labels.font.size = Pt(11)
pl.series[0].format.fill.solid(); pl.series[0].format.fill.fore_color.rgb = GREEN
table(sl, 8.1, 1.7, 4.6, ["비목", "비율"], [(b[0], f"{b[2] / tot * 100:.0f}%") for b in D.BUDGET] + [("합계 5,000만 원", "100%")], [3.2, 1.4], 12, 0.5)

# 23 타당성
sl = base("타당성: 조건부 가능 — 기존 설비 연동과 작물 정식 시점이 관건", notes="결론은 조건부 가능입니다. 기존 설비를 쓸 수 있어야 하고, 정식이 2월이라 1월 인수 시점에는 작물이 없으므로 실작물 검증은 인수 후 1작기 운영에서 하겠습니다.")
table(sl, 0.6, 1.6, 12.1, ["원래 계획", "검토 결과", "수정 권고", "영향"], [
    ("실장비 + 웹을 5,000만 원에", "기존 설비 연동이면 가능", "연동을 계약 조건으로", "미충족 시 F-03·04 축소"),
    ("인건비 2,500만 원", "공표 노임단가 적용 시 초과 가능", "가상 단가임을 명시", "실사업은 증액·축소"),
    ("12주에 기능 18건", "구현 5주로 빠듯", "MVP 16건 우선", "확장 2건 미완료 가능"),
    ("기간 내 작물 검증", "1월 인수 시 작물 없음", "인수 후 1작기 운영 검증", "KPI 일부 인수 후 측정"),
    ("겨울 설치", "한파 지연 위험", "설치 3주, 실내 사전 조립", "일정 위험 중간")], [2.8, 3, 3, 3.3], 12.5, 0.62)

# 24 위험
sl = base("위험 8건 중 최우선은 기존 설비 연동(R-01)", notes="가능성과 영향을 곱해 점수를 매겼습니다. 12점 이상은 R-01 하나이고, 착수 첫 주 현장 조사로 바로 확인하겠습니다.")
fit_pic(sl, "f_risk_matrix", 0.6, 1.4, 5.6, 5.4)
top = sorted(D.RISKS, key=lambda r: -r[2] * r[3])[:5]
table(sl, 6.5, 1.6, 6.2, ["ID", "위험", "점수", "대응"], [(r[0], r[1], r[2] * r[3], r[4]) for r in top], [0.9, 2.4, 0.8, 3.4], 11, 0.75)

# 25 테스트·KPI
sl = base("테스트 19건·KPI 10개는 모두 '미실행 / 측정 대기'", "보고서에서 시뮬레이션 검증(B)과 실장비 검증(A)을 구분해 채운다", notes="아직 개발 전이므로 결과는 하나도 적지 않았습니다. 개발 보고서와 완료 보고서에서 실제 로그와 측정값으로만 채우겠습니다.")
table(sl, 0.6, 1.7, 6.4, ["KPI", "목표"], [(k[0], k[1]) for k in D.KPIS[:7]], [3, 3.4], 11.5, 0.55)
table(sl, 7.3, 1.7, 5.4, ["TC", "합격 기준", "결과"], [(t[0], t[3], "미실행") for t in D.TESTS[:8]], [1.1, 3.2, 1.1], 11, 0.5)

# 26 AI·확인 필요
sl = base("남은 확인 사항과 AI 3종 비교 계획", notes="교수님께 확인받아야 할 것은 지정 농업 DB 사이트입니다. 재배 기준은 과제 조건대로 네이버로 다시 확인하겠습니다. AI 비교는 같은 질문으로 직접 실행한 뒤 평가표를 채웁니다.")
card(sl, 0.6, 1.6, 5.9, 5.1)
text(sl, 0.85, 1.75, 5.4, 4.8, ["**확인 필요**", "B-1 교수 지정 농업 DB 명칭·URL", "B-2 네이버로 재배 기준값 재확인", "B-3 양평 ASOS 지점 202·결측률", "B-4 하우스 규격·재배 방식·품종·작형", "B-5 기존 설비 연동 가능 여부"], 14)
table(sl, 6.8, 1.6, 5.9, ["평가 기준", "ChatGPT", "Claude", "Gemini"], [(k, "평가 대기", "평가 대기", "평가 대기") for k in ("정확성", "구체성", "실현 가능성", "출처 신뢰성", "환각 가능성")],
      [2, 1.3, 1.3, 1.3], 11, 0.55)
text(sl, 6.8, 5.2, 5.9, 1.4, "같은 질문을 세 AI에 새 대화로 입력해 비교한다. 실행하지 않은 결과는 채우지 않는다.", 12, False, MUTED)

# 27 결론
sl = base("", dark=True, notes="정리하겠습니다. 교수님 조건 안에서 MVP 16건은 구현 가능합니다. 오늘 승인받고 싶은 것은 요구사항 범위, 2구역 배치안, 그리고 실장비 검증을 인수 후로 넘기는 일정 조정입니다.")
text(sl, 0.8, 1.0, 11.5, 0.9, "결론: 조건 안에서 MVP 16건은 구현 가능", 32, True, WHITE)
for i, (h, b) in enumerate([("승인 요청", "요구사항 25건 범위\n2구역 배치안\n실작물 검증은 인수 후"), ("다음 단계", "착수 1주 현장 조사\nB-1~B-5 확인\nG1 요구사항 확정 (2주차)"), ("산출물", "계획서 Word·PPT\nRFP 초안\nDB 스키마 20개")]):
    card(sl, 0.8 + i * 4.0, 2.6, 3.7, 3.4, fill=RGBColor(0x26, 0x50, 0x2B))
    text(sl, 1.05 + i * 4.0, 2.8, 3.2, 0.5, h, 18, True, MOSS)
    text(sl, 1.05 + i * 4.0, 3.45, 3.2, 2.4, b, 15, False, WHITE)

prs.save(OUT)
print(OUT, len(prs.slides), "slides")
